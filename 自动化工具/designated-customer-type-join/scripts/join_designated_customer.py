#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
指定客户类型关联 —— 门店定义数据加工步骤。

作用:
  把「指定客户城市总部.xlsx」(依赖文件 B) 的「指定客户类型」，
  按客户编码关联进「生产门店数据_By条件汇总」(输入文件 A) 的最后一列，
  结果写成独立新文件，不覆盖原输入文件。

关联键 (已实测验证):
  A 的「二级城市总部编码」 = B 的「客户编码」

  注意 —— 口语里的「全国总部编码」在文件里就是「二级城市总部编码」这一列：
    - B 的「客户属性」虽然全是「全国总部」，但「客户编码」落在 A 的二级城市总部这一层；
    - A 的「一级/三级城市总部编码」与 B 的客户编码 0 匹配，用了会整列 #N/A。

用法:
  python join_designated_customer.py                  # 自动取输入目录下日期最新的一份
  python join_designated_customer.py --input <A> --deps <B> --output-dir <dir>
"""

import argparse
import glob
import re
from pathlib import Path

import openpyxl

# 默认值 = 当前 2027 门店定义数据加工任务的真实路径与列名
DEFAULTS = {
    "input_dir": r"D:\CMJWORK\003-CIP\2027需求\数据处理\输入文件",
    "deps": r"D:\CMJWORK\003-CIP\2027需求\数据处理\依赖文件\指定客户城市总部.xlsx",
    "output_dir": r"D:\CMJWORK\003-CIP\2027需求\数据处理\输出文件",
    "sheet_a": "按条件统计门店数",
    "sheet_b": "sheet1",  # B 的 sheet 名是小写 sheet1
    "join_col_a": "二级城市总部编码",
    "join_col_b": "客户编码",
    "value_col": "指定客户类型",
    "target_col": "指定客户类型",  # 写入 A 的列名；不存在则追加为新列
    "unmatched": "",  # 未匹配到的行填什么（默认留空）
}


def norm(v):
    """统一键值类型：去空白、None 转空。避免数字/字符串不一致导致匹配不上。"""
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def header_index(header_row, name):
    """在表头元组里找列名，返回 0-based 索引；找不到返回 None。"""
    for i, h in enumerate(header_row):
        if h is not None and str(h).strip() == name:
            return i
    return None


def find_latest_input(input_dir):
    """在输入目录里找 生产门店数据_By条件汇总_<日期>.xlsx，返回日期最大的一份路径。

    按文件名尾部的日期数字（如 1024）从大到小排序，而不是按文件修改时间，
    这样即使旧日期文件被重新导出、mtime 更新，也不会被误选。
    """
    pattern = str(input_dir / "生产门店数据_By条件汇总_*.xlsx")
    candidates = []
    for p in glob.glob(pattern):
        m = re.search(r"生产门店数据_By条件汇总_(\d+)\.xlsx$", p)
        if m:
            candidates.append((int(m.group(1)), p))
    if not candidates:
        return None
    candidates.sort(key=lambda x: x[0], reverse=True)
    return Path(candidates[0][1])


def build_b_map(deps_path, sheet_b, join_col_b, value_col):
    """读依赖文件 B，返回 {客户编码: 指定客户类型}。"""
    wb = openpyxl.load_workbook(deps_path, read_only=True, data_only=True)
    if sheet_b not in wb.sheetnames:
        wb.close()
        raise KeyError(f"依赖文件里没有 sheet「{sheet_b}」，实际有 {wb.sheetnames}")
    ws = wb[sheet_b]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("依赖文件 B 为空")

    join_i = header_index(rows[0], join_col_b)
    value_i = header_index(rows[0], value_col)
    if join_i is None:
        raise KeyError(f"B 的表头里没有列「{join_col_b}」，实际有 {rows[0]}")
    if value_i is None:
        raise KeyError(f"B 的表头里没有列「{value_col}」，实际有 {rows[0]}")

    b_map = {}
    for r in rows[1:]:
        key = norm(r[join_i])
        val = norm(r[value_i])
        if key and val:  # 客户编码 或 指定客户类型 为空的记录跳过
            b_map[key] = val
    return b_map


def main():
    p = argparse.ArgumentParser(
        description="把「指定客户类型」按客户编码关联进生产门店数据，输出独立新文件。"
    )
    p.add_argument("--input", default=None, help="输入文件 A 路径（缺省自动取输入目录下最新一份）")
    p.add_argument("--input-dir", default=DEFAULTS["input_dir"], help="输入目录（--input 缺省时在此自动找最新）")
    p.add_argument("--deps", default=DEFAULTS["deps"], help="依赖文件 B 路径")
    p.add_argument("--output-dir", default=DEFAULTS["output_dir"], help="输出目录")
    p.add_argument("--output-name", default=None, help="输出文件名（默认与输入同名）")
    p.add_argument("--sheet-a", default=DEFAULTS["sheet_a"], help="A 里要处理的工作表")
    p.add_argument("--sheet-b", default=DEFAULTS["sheet_b"], help="B 里读关联键的工作表")
    p.add_argument("--join-col-a", default=DEFAULTS["join_col_a"], help="A 的关联键列名")
    p.add_argument("--join-col-b", default=DEFAULTS["join_col_b"], help="B 的关联键列名")
    p.add_argument("--value-col", default=DEFAULTS["value_col"], help="B 中要写入的列名")
    p.add_argument("--target-col", default=DEFAULTS["target_col"], help="写入 A 的列名")
    p.add_argument("--unmatched", default=DEFAULTS["unmatched"], help="未匹配行填的值")
    args = p.parse_args()

    if args.input:
        input_path = Path(args.input)
    else:
        input_path = find_latest_input(Path(args.input_dir))
        if input_path is None:
            raise FileNotFoundError(
                f"输入目录里没有 生产门店数据_By条件汇总_*.xlsx: {args.input_dir}"
            )
        print(f"[输入] 自动取最新: {input_path}")
    deps_path = Path(args.deps)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    if not input_path.exists():
        raise FileNotFoundError(f"输入文件不存在: {input_path}")
    if not deps_path.exists():
        raise FileNotFoundError(f"依赖文件不存在: {deps_path}")

    # 1) 读 B 建映射
    b_map = build_b_map(deps_path, args.sheet_b, args.join_col_b, args.value_col)
    print(f"[B] {args.sheet_b}: 读到 {len(b_map)} 条 客户编码 -> 指定客户类型")

    # 2) 读 A（保留原格式与公式，非 read_only；data_only=False 保留其它列公式）
    wb_a = openpyxl.load_workbook(input_path, data_only=False)
    if args.sheet_a not in wb_a.sheetnames:
        wb_a.close()
        raise KeyError(f"输入文件里没有 sheet「{args.sheet_a}」，实际有 {wb_a.sheetnames}")
    ws_a = wb_a[args.sheet_a]

    header = [c.value for c in ws_a[1]]
    join_a_i = header_index(tuple(header), args.join_col_a)
    if join_a_i is None:
        wb_a.close()
        raise KeyError(f"A 的表头里没有列「{args.join_col_a}」，实际有 {header}")

    # 目标列：已存在则复用，否则追加为新列
    target_i = header_index(tuple(header), args.target_col)
    if target_i is None:
        target_i = ws_a.max_column  # 0-based -> 追加到最后
        ws_a.cell(row=1, column=target_i + 1, value=args.target_col)
        print(f"[A] 未找到「{args.target_col}」列，已追加为新列")
    else:
        print(f"[A] 复用已有「{args.target_col}」列（原 #N/A 将被覆盖）")

    # 3) 逐行关联写入
    matched = unmatched = 0
    samples = []
    for row in range(2, ws_a.max_row + 1):
        key = norm(ws_a.cell(row=row, column=join_a_i + 1).value)
        if key is not None and key in b_map:
            ws_a.cell(row=row, column=target_i + 1).value = b_map[key]
            matched += 1
            if len(samples) < 5:
                samples.append((key, b_map[key]))
        else:
            ws_a.cell(row=row, column=target_i + 1).value = args.unmatched
            unmatched += 1

    # 4) 保存为独立新文件
    output_path = output_dir / (args.output_name or input_path.name)
    wb_a.save(output_path)
    wb_a.close()

    total = matched + unmatched
    print(f"[A] {args.sheet_a}: 共 {total} 行，匹配 {matched}，未匹配 {unmatched}")
    print(f"样例匹配: {samples}")
    print(f"已输出: {output_path}")


if __name__ == "__main__":
    main()
